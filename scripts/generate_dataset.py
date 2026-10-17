"""
Customer Segmentation & Retention Analytics - Synthetic Dataset Generator
Project: Customer Segmentation & Retention Analytics in Power BI
Workspace: D:/ML/Sentimental Analysis

Generates a reproducible 150-row synthetic customer dataset with realistic
Indian customer profiles, purchasing behavior, and deterministic business rules.
Outputs:
  - data/raw/customer_segmentation_retention_sample.csv
  - data/raw/customer_segmentation_retention_sample.xlsx (Sheets: 'Customer Dataset', 'Data Dictionary')
"""

import os
import random
import datetime
from decimal import Decimal, ROUND_HALF_UP
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Explicit Random Seed for Strict Reproducibility
RANDOM_SEED = 42
SNAPSHOT_DATE = datetime.date(2026, 10, 16)
NUM_RECORDS = 150

FIRST_NAMES = [
    "Aarav", "Vivaan", "Aditya", "Vihaan", "Arjun", "Sai", "Reyansh", "Ayaan", "Krishna", "Ishaan",
    "Shaurya", "Atharva", "Dhruv", "Kabir", "Rohan", "Ananya", "Diya", "Isha", "Riya", "Aadhya",
    "Kavya", "Pooja", "Priya", "Neha", "Sneha", "Tanvi", "Sanya", "Meera", "Anika", "Tara",
    "Vikram", "Rahul", "Amit", "Rajesh", "Suresh", "Sunita", "Deepa", "Shweta", "Karan", "Nikhil",
    "Preeti", "Sanjay", "Alok", "Gaurav", "Manoj", "Simran", "Swati", "Ritu", "Harish", "Varun"
]

LAST_NAMES = [
    "Sharma", "Verma", "Patel", "Gupta", "Mehta", "Reddy", "Nair", "Iyer", "Rao", "Kumar",
    "Singh", "Choudhury", "Joshi", "Bhat", "Deshmukh", "Kulkarni", "Chatterjee", "Banerjee", "Das", "Sengupta",
    "Kapoor", "Malhotra", "Bhasin", "Saxena", "Agarwal", "Mishra", "Pandey", "Trivedi", "Chopra", "Dutta"
]

CITIES = [
    "Mumbai, Maharashtra", "Delhi, NCR", "Bengaluru, Karnataka", "Hyderabad, Telangana",
    "Chennai, Tamil Nadu", "Kolkata, West Bengal", "Pune, Maharashtra", "Ahmedabad, Gujarat",
    "Jaipur, Rajasthan", "Surat, Gujarat", "Lucknow, Uttar Pradesh", "Kochi, Kerala",
    "Chandigarh, Punjab", "Indore, Madhya Pradesh", "Nagpur, Maharashtra", "Bhopal, Madhya Pradesh"
]

OCCUPATIONS = [
    "Software Engineer", "Business Owner", "Data Analyst", "Chartered Accountant",
    "Marketing Manager", "Healthcare Professional", "Educator / Professor", "Financial Advisor",
    "Retail Store Manager", "Consultant", "Civil Engineer", "Architect", "Research Scientist", "Government Officer"
]

INCOME_LEVELS = [
    "Low (< ₹5L)",
    "Lower-Middle (₹5L - ₹10L)",
    "Upper-Middle (₹10L - ₹25L)",
    "High (> ₹25L)"
]

PRODUCT_CATEGORIES = [
    "Electronics & Gadgets",
    "Fashion & Apparel",
    "Home & Kitchen",
    "Beauty & Personal Care",
    "Sports & Fitness",
    "Books & Media",
    "Groceries & Essentials"
]

PURCHASE_CHANNELS = [
    "Mobile App",
    "Website",
    "In-Store",
    "Omnichannel"
]

def calculate_retention_status(last_purchase_date: datetime.date, snapshot_date: datetime.date) -> str:
    """Calculates Retention Status based on elapsed days from snapshot."""
    days = (snapshot_date - last_purchase_date).days
    if days <= 60:
        return "Retained"
    elif days <= 150:
        return "At Risk"
    else:
        return "Churned"

def calculate_loyalty_tier(total_spend: float) -> str:
    """Calculates Loyalty Tier based on total spend."""
    if total_spend < 15000.00:
        return "Bronze"
    elif total_spend < 40000.00:
        return "Silver"
    elif total_spend < 80000.00:
        return "Gold"
    else:
        return "Platinum"

def calculate_customer_segment(retention_status: str, loyalty_tier: str, purchase_frequency: int, customer_type: str) -> str:
    """
    Deterministic, mutually exclusive Customer Segment assignment.
    Every customer maps to exactly one segment:
      - Champions
      - Loyal Customers
      - New / Potential Loyalist
      - At Risk - High Value
      - Needs Nurturing
      - Hibernating
    """
    if retention_status == "Retained":
        if loyalty_tier in ["Gold", "Platinum"] and purchase_frequency >= 5:
            return "Champions"
        elif loyalty_tier in ["Silver", "Gold", "Platinum"]:
            return "Loyal Customers"
        else:
            return "New / Potential Loyalist"
    elif retention_status == "At Risk":
        if loyalty_tier in ["Gold", "Platinum"]:
            return "At Risk - High Value"
        elif customer_type == "New":
            return "New / Potential Loyalist"
        else:
            return "Needs Nurturing"
    else:  # Churned (> 150 days)
        if loyalty_tier in ["Gold", "Platinum"]:
            return "At Risk - High Value"
        else:
            return "Hibernating"

def generate_dataset(seed: int = RANDOM_SEED) -> pd.DataFrame:
    """Generates the 150-record synthetic dataset deterministically."""
    random.seed(seed)
    records = []

    # Target distributions across archetypes to ensure all segments, tiers, and statuses are populated
    # Archetype distribution guides:
    # 0: Champions target
    # 1: Loyal Customers target
    # 2: New / Potential Loyalist target
    # 3: At Risk - High Value target
    # 4: Needs Nurturing target
    # 5: Hibernating target

    archetype_distribution = [0] * 32 + [1] * 28 + [2] * 25 + [3] * 22 + [4] * 23 + [5] * 20
    random.shuffle(archetype_distribution)

    used_names = set()

    for i in range(NUM_RECORDS):
        cust_id = f"IND-{1001 + i}"
        archetype = archetype_distribution[i]

        # Generate unique fictional name
        while True:
            fn = random.choice(FIRST_NAMES)
            ln = random.choice(LAST_NAMES)
            name = f"{fn} {ln}"
            if name not in used_names:
                used_names.add(name)
                break

        gender = random.choice(["Male", "Female", "Other"])
        city = random.choice(CITIES)
        occupation = random.choice(OCCUPATIONS)
        channel = random.choice(PURCHASE_CHANNELS)
        category = random.choice(PRODUCT_CATEGORIES)

        # Behavioral attributes conditioned on archetype to guarantee rich, realistic variation
        if archetype == 0:  # Champions
            age = random.randint(28, 58)
            income = random.choice(["Upper-Middle (₹10L - ₹25L)", "High (> ₹25L)"])
            # Retained (0-60 days)
            days_recency = random.randint(2, 50)
            purchase_freq = random.randint(5, 18)
            # Tier: Gold or Platinum (>= 40,000)
            aov = round(random.uniform(5000.0, 14000.0), 2)
            # Ensure total spend >= 40,000
            if purchase_freq * aov < 40000.0:
                aov = round(42000.0 / purchase_freq + random.uniform(500.0, 2000.0), 2)
            engagement = random.randint(75, 98)
            satisfaction = round(random.uniform(4.2, 5.0), 1)

        elif archetype == 1:  # Loyal Customers
            age = random.randint(25, 62)
            income = random.choice(["Lower-Middle (₹5L - ₹10L)", "Upper-Middle (₹10L - ₹25L)", "High (> ₹25L)"])
            days_recency = random.randint(5, 58)
            # Freq can be 2-4 for Gold/Platinum or >=2 for Silver
            if random.random() < 0.6:
                # Silver tier
                purchase_freq = random.randint(2, 7)
                aov = round(random.uniform(3000.0, 7500.0), 2)
                # Ensure 15,000 <= spend < 40,000
                total_target = random.uniform(16000.0, 38000.0)
                aov = round(total_target / purchase_freq, 2)
            else:
                # Gold/Platinum with frequency < 5
                purchase_freq = random.randint(2, 4)
                total_target = random.uniform(41000.0, 85000.0)
                aov = round(total_target / purchase_freq, 2)
            engagement = random.randint(60, 88)
            satisfaction = round(random.uniform(3.8, 4.8), 1)

        elif archetype == 2:  # New / Potential Loyalist
            age = random.randint(20, 50)
            income = random.choice(["Low (< ₹5L)", "Lower-Middle (₹5L - ₹10L)", "Upper-Middle (₹10L - ₹25L)"])
            is_new = random.random() < 0.65
            if is_new:
                purchase_freq = 1
                days_recency = random.randint(3, 140)  # Retained or At Risk
                aov = round(random.uniform(1200.0, 14500.0), 2)
            else:
                # Returning, Retained, Bronze tier (< 15,000)
                purchase_freq = random.randint(2, 4)
                days_recency = random.randint(4, 55)
                aov = round(random.uniform(1500.0, 3200.0), 2)
                if purchase_freq * aov >= 15000.0:
                    aov = round(14000.0 / purchase_freq, 2)
            engagement = random.randint(50, 85)
            satisfaction = round(random.uniform(3.5, 4.6), 1)

        elif archetype == 3:  # At Risk - High Value
            age = random.randint(30, 68)
            income = random.choice(["Upper-Middle (₹10L - ₹25L)", "High (> ₹25L)"])
            # At Risk (61-150) or Churned (> 150)
            days_recency = random.choice([random.randint(65, 145), random.randint(155, 310)])
            purchase_freq = random.randint(3, 12)
            # Gold or Platinum tier (>= 40,000)
            total_target = random.uniform(42000.0, 115000.0)
            aov = round(total_target / purchase_freq, 2)
            engagement = random.randint(30, 65)
            satisfaction = round(random.uniform(2.5, 3.9), 1)

        elif archetype == 4:  # Needs Nurturing
            age = random.randint(22, 60)
            income = random.choice(["Low (< ₹5L)", "Lower-Middle (₹5L - ₹10L)", "Upper-Middle (₹10L - ₹25L)"])
            # At Risk (61-150 days), Returning, Bronze or Silver (< 40,000)
            days_recency = random.randint(62, 148)
            purchase_freq = random.randint(2, 6)
            total_target = random.uniform(5000.0, 38000.0)
            aov = round(total_target / purchase_freq, 2)
            engagement = random.randint(25, 60)
            satisfaction = round(random.uniform(2.2, 3.8), 1)

        else:  # Hibernating
            age = random.randint(24, 75)
            income = random.choice(["Low (< ₹5L)", "Lower-Middle (₹5L - ₹10L)", "Upper-Middle (₹10L - ₹25L)"])
            # Churned (> 150 days), Bronze or Silver (< 40,000)
            days_recency = random.randint(155, 420)
            purchase_freq = random.randint(1, 4)
            total_target = random.uniform(1500.0, 35000.0)
            aov = round(total_target / purchase_freq, 2)
            engagement = random.randint(5, 35)
            satisfaction = round(random.uniform(1.2, 3.2), 1)

        # Compute Dates
        last_purchase_date = SNAPSHOT_DATE - datetime.timedelta(days=days_recency)
        # Join Date must be <= Last Purchase Date
        # Customer tenure before last purchase
        if purchase_freq == 1:
            tenure_before_last_purchase = random.randint(0, 14)  # bought shortly after joining
        else:
            tenure_before_last_purchase = random.randint(15, 600)
        join_date = last_purchase_date - datetime.timedelta(days=tenure_before_last_purchase)

        # Business calculations
        # Ensure Total Spend = Purchase Frequency * Average Order Value with standard 2 decimal places
        total_spend = round(purchase_freq * aov, 2)
        customer_type = "New" if purchase_freq == 1 else "Returning"
        retention_status = calculate_retention_status(last_purchase_date, SNAPSHOT_DATE)
        loyalty_tier = calculate_loyalty_tier(total_spend)
        customer_segment = calculate_customer_segment(retention_status, loyalty_tier, purchase_freq, customer_type)

        records.append({
            "Customer ID": cust_id,
            "Customer Name": name,
            "Age": age,
            "Gender": gender,
            "City/Region": city,
            "Occupation": occupation,
            "Income Level": income,
            "Join Date": join_date.strftime("%Y-%m-%d"),
            "Last Purchase Date": last_purchase_date.strftime("%Y-%m-%d"),
            "Purchase Frequency": purchase_freq,
            "Average Order Value": aov,
            "Total Spend": total_spend,
            "Preferred Product Category": category,
            "Purchase Channel": channel,
            "Customer Type": customer_type,
            "Engagement Score": engagement,
            "Satisfaction Score": satisfaction,
            "Retention Status": retention_status,
            "Customer Segment": customer_segment,
            "Loyalty Tier": loyalty_tier
        })

    df = pd.DataFrame(records)
    return df

DATA_DICTIONARY_ROWS = [
    {
        "Column Name": "Customer ID",
        "Data Type": "String (Text)",
        "Nullable": "No",
        "Business Definition": "Unique fictional identifier assigned to each customer record (format: IND-xxxx).",
        "Validation Rule": "Must be unique, non-empty, matching regex ^IND-\\d{4}$."
    },
    {
        "Column Name": "Customer Name",
        "Data Type": "String (Text)",
        "Nullable": "No",
        "Business Definition": "Fictional full name of the customer, synthesized for privacy without real PII.",
        "Validation Rule": "Non-empty string containing first and last name."
    },
    {
        "Column Name": "Age",
        "Data Type": "Integer",
        "Nullable": "No",
        "Business Definition": "Customer age in complete years.",
        "Validation Rule": "Must be an integer in range [18, 90]."
    },
    {
        "Column Name": "Gender",
        "Data Type": "String (Categorical)",
        "Nullable": "No",
        "Business Definition": "Reported customer gender identity.",
        "Validation Rule": "Must be one of ['Male', 'Female', 'Other']."
    },
    {
        "Column Name": "City/Region",
        "Data Type": "String (Text)",
        "Nullable": "No",
        "Business Definition": "Customer geographical location denominated by Indian City and State.",
        "Validation Rule": "Must be a non-empty string in valid 'City, State' format."
    },
    {
        "Column Name": "Occupation",
        "Data Type": "String (Categorical)",
        "Nullable": "No",
        "Business Definition": "Primary professional occupation of the customer.",
        "Validation Rule": "Non-empty recognized occupational category."
    },
    {
        "Column Name": "Income Level",
        "Data Type": "String (Categorical)",
        "Nullable": "No",
        "Business Definition": "Self-reported annual household income bracket in Indian Rupees (INR).",
        "Validation Rule": "Must be one of ['Low (< ₹5L)', 'Lower-Middle (₹5L - ₹10L)', 'Upper-Middle (₹10L - ₹25L)', 'High (> ₹25L)']."
    },
    {
        "Column Name": "Join Date",
        "Data Type": "Date (YYYY-MM-DD)",
        "Nullable": "No",
        "Business Definition": "Date customer registered or created their primary account.",
        "Validation Rule": "Valid ISO date; must be <= Last Purchase Date."
    },
    {
        "Column Name": "Last Purchase Date",
        "Data Type": "Date (YYYY-MM-DD)",
        "Nullable": "No",
        "Business Definition": "Date of the most recent completed transaction.",
        "Validation Rule": "Valid ISO date; must be >= Join Date and <= 2026-10-16 (Snapshot Date)."
    },
    {
        "Column Name": "Purchase Frequency",
        "Data Type": "Integer",
        "Nullable": "No",
        "Business Definition": "Lifetime count of distinct purchases completed by the customer.",
        "Validation Rule": "Positive integer (>= 1)."
    },
    {
        "Column Name": "Average Order Value",
        "Data Type": "Decimal / Currency (INR)",
        "Nullable": "No",
        "Business Definition": "Mean gross purchase amount per order denominated in INR (₹).",
        "Validation Rule": "Non-negative float rounded to 2 decimal places."
    },
    {
        "Column Name": "Total Spend",
        "Data Type": "Decimal / Currency (INR)",
        "Nullable": "No",
        "Business Definition": "Lifetime cumulative monetary spend across all orders in INR (₹).",
        "Validation Rule": "Must equal Purchase Frequency * Average Order Value (+/- 0.01 rounding tolerance)."
    },
    {
        "Column Name": "Preferred Product Category",
        "Data Type": "String (Categorical)",
        "Nullable": "No",
        "Business Definition": "Merchandise category with the highest frequency of purchase.",
        "Validation Rule": "Must be one of the 7 designated retail merchandise categories."
    },
    {
        "Column Name": "Purchase Channel",
        "Data Type": "String (Categorical)",
        "Nullable": "No",
        "Business Definition": "Primary retail sales channel utilized by the customer.",
        "Validation Rule": "Must be one of ['Mobile App', 'Website', 'In-Store', 'Omnichannel']."
    },
    {
        "Column Name": "Customer Type",
        "Data Type": "String (Categorical)",
        "Nullable": "No",
        "Business Definition": "Classification indicating purchase lifecycle stage (New vs. Returning).",
        "Validation Rule": "Must be 'New' if Purchase Frequency == 1, otherwise 'Returning'."
    },
    {
        "Column Name": "Engagement Score",
        "Data Type": "Integer",
        "Nullable": "No",
        "Business Definition": "Composite behavioral score indexing digital interactions, visits, and clicks.",
        "Validation Rule": "Integer in range [1, 100]."
    },
    {
        "Column Name": "Satisfaction Score",
        "Data Type": "Decimal",
        "Nullable": "No",
        "Business Definition": "Customer satisfaction rating aggregated from feedback surveys (CSAT).",
        "Validation Rule": "Decimal rating in range [1.0, 5.0]."
    },
    {
        "Column Name": "Retention Status",
        "Data Type": "String (Categorical)",
        "Nullable": "No",
        "Business Definition": "Active retention state calculated based on recency from snapshot 2026-10-16.",
        "Validation Rule": "'Retained' (0-60d), 'At Risk' (61-150d), or 'Churned' (>150d)."
    },
    {
        "Column Name": "Customer Segment",
        "Data Type": "String (Categorical)",
        "Nullable": "No",
        "Business Definition": "Mutually exclusive RFM behavioral segment for targeted marketing and retention.",
        "Validation Rule": "Must belong to exactly one of: ['Champions', 'Loyal Customers', 'New / Potential Loyalist', 'At Risk - High Value', 'Needs Nurturing', 'Hibernating']."
    },
    {
        "Column Name": "Loyalty Tier",
        "Data Type": "String (Categorical)",
        "Nullable": "No",
        "Business Definition": "Spend-based rewards tier determining perks and prioritization.",
        "Validation Rule": "'Bronze' (<15k), 'Silver' (15k-40k), 'Gold' (40k-80k), 'Platinum' (>=80k INR)."
    }
]

def export_to_excel(df: pd.DataFrame, output_path: str):
    """Exports dataset and formatted data dictionary to a clean multi-sheet Excel workbook."""
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # 1. Customer Dataset Sheet
    ws_data = wb.create_sheet(title="Customer Dataset")
    ws_data.views.sheetView[0].showGridLines = True

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    data_font = Font(name="Segoe UI", size=10)
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    headers = list(df.columns)
    ws_data.append(headers)

    for col_num in range(1, len(headers) + 1):
        cell = ws_data.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, row in df.iterrows():
        row_values = list(row)
        ws_data.append(row_values)
        row_num = r_idx + 2
        for col_num in range(1, len(headers) + 1):
            cell = ws_data.cell(row=row_num, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            # Numeric alignment and formatting
            col_name = headers[col_num - 1]
            if col_name in ["Average Order Value", "Total Spend"]:
                cell.number_format = '₹#,##0.00'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_name in ["Purchase Frequency", "Age", "Engagement Score"]:
                cell.number_format = '#,##0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_name == "Satisfaction Score":
                cell.number_format = '0.0'
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif "Date" in col_name or col_name == "Customer ID":
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

    # Auto-adjust column widths
    for col in ws_data.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_data.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # 2. Data Dictionary Sheet
    ws_dict = wb.create_sheet(title="Data Dictionary")
    ws_dict.views.sheetView[0].showGridLines = True

    dict_headers = ["Column Name", "Data Type", "Nullable", "Business Definition", "Validation Rule"]
    ws_dict.append(dict_headers)

    for col_num in range(1, len(dict_headers) + 1):
        cell = ws_dict.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r_idx, row_dict in enumerate(DATA_DICTIONARY_ROWS):
        vals = [row_dict[h] for h in dict_headers]
        ws_dict.append(vals)
        row_num = r_idx + 2
        for col_num in range(1, len(dict_headers) + 1):
            cell = ws_dict.cell(row=row_num, column=col_num)
            cell.font = data_font
            cell.border = thin_border
            if col_num in [1, 2, 3]:
                cell.alignment = Alignment(horizontal="center", vertical="top")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    for col in ws_dict.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_dict.column_dimensions[col_letter].width = min(max(max_len + 4, 15), 55)

    wb.save(output_path)

def main():
    # Resolve relative paths from project root
    project_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    raw_data_dir = os.path.join(project_root, "data", "raw")
    os.makedirs(raw_data_dir, exist_ok=True)

    csv_path = os.path.join(raw_data_dir, "customer_segmentation_retention_sample.csv")
    xlsx_path = os.path.join(raw_data_dir, "customer_segmentation_retention_sample.xlsx")

    print(f"Generating synthetic dataset with seed {RANDOM_SEED} (Snapshot Date: {SNAPSHOT_DATE})...")
    df = generate_dataset(seed=RANDOM_SEED)

    print(f"Exporting CSV: {csv_path}")
    df.to_csv(csv_path, index=False, encoding="utf-8")

    print(f"Exporting Excel Workbook: {xlsx_path}")
    export_to_excel(df, xlsx_path)

    print(f"SUCCESS: Generated {len(df)} records across {len(df.columns)} columns.")
    print("Retention Breakdown:\n", df["Retention Status"].value_counts())
    print("\nCustomer Segment Breakdown:\n", df["Customer Segment"].value_counts())
    print("\nLoyalty Tier Breakdown:\n", df["Loyalty Tier"].value_counts())

if __name__ == "__main__":
    main()
